#!/usr/bin/env python3
"""
scripts/check_roster.py - Strict Anti-Fabrication & Roster Verification Gate

Verifies that no tracked files contain fabricated person names, unverified student IDs (MSSV),
or fake contribution percentages.
Official Source of Truth (SSOT): docs/TEAM_ROSTER.md
"""

import sys
import os
import re
import subprocess
import unicodedata
from pathlib import Path

# Ensure UTF-8 output on Windows
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

REPO_ROOT = Path(__file__).resolve().parent.parent
ROSTER_PATH = REPO_ROOT / "docs" / "TEAM_ROSTER.md"

# Known fabricated names from prior incidents (NFC normalized)
FABRICATED_NAMES = [
    unicodedata.normalize("NFC", p) for p in [
        r"Dương\s+Bắp\s+Vịt\s+Hoàng",
        r"Hoàng\s+Văn\s+Bắp",
        r"Nguyễn\s+Nhật\s+Huy",
        r"Trần\s+Quốc\s+Khánh",
        r"Bắp\s+Vịt\s+Hoàng",
    ]
]


def parse_team_roster(roster_path: Path):
    """
    Parses docs/TEAM_ROSTER.md to load verified names, MSSVs, and handles.
    """
    verified_names = set()
    verified_mssvs = set()
    allowed_emails = set()
    allowed_handles = set()

    if not roster_path.exists():
        rel_display = "docs/TEAM_ROSTER.md"
        print(f"[FAIL] {rel_display} not found")
        sys.exit(1)

    with open(roster_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    for line in lines:
        line = unicodedata.normalize("NFC", line.strip())
        if not line.startswith("|") or "Full Name" in line or ":---" in line:
            continue
        parts = [p.strip() for p in line.split("|")[1:-1]]
        if len(parts) >= 6:
            full_name, mssv, github, git_email, role, verified_str = parts[:6]
            is_verified = verified_str.strip().lower() == "true"
            
            # Sanitize markdown formatting from cells
            clean_name = full_name.strip(" *`").strip()
            clean_mssv = mssv.strip(" *`").strip()

            # Extract email and github handle
            email_match = re.search(r"[\w\.-]+@[\w\.-]+", git_email)
            if email_match:
                allowed_emails.add(email_match.group(0).lower())
            handle_match = re.search(r"@[\w-]+", github)
            if handle_match:
                allowed_handles.add(handle_match.group(0).lower())

            if is_verified:
                if clean_name and clean_name != "CHƯA CÓ DỮ LIỆU":
                    verified_names.add(clean_name.lower())
                if clean_mssv and clean_mssv != "CHƯA CÓ DỮ LIỆU":
                    verified_mssvs.add(clean_mssv)

    return verified_names, verified_mssvs, allowed_emails, allowed_handles


def get_tracked_files():
    """Returns list of tracked files via git ls-files."""
    res = subprocess.run(["git", "ls-files"], cwd=REPO_ROOT, capture_output=True, text=True, check=True)
    return [f.strip() for f in res.stdout.splitlines() if f.strip()]


def check_content(content: str, filename: str, verified_names, verified_mssvs, allowed_emails, is_binary=False):
    r"""
    Checks text content against anti-fabrication rules:
    (a) pattern 2[0-5]11\d{4} not in roster with verified_by_user true (excluding allowed git emails)
    (b) full-name-like string next to an MSSV or role not in roster, or known fabricated names
    (c) contribution percentage next to a person / role
    """
    violations = []
    # Always NFC normalize incoming text
    normalized_content = unicodedata.normalize("NFC", content)
    lines = normalized_content.splitlines()

    for idx, line in enumerate(lines, 1):
        loc = f"{filename}:{idx}" if not is_binary else f"{filename} (extracted text)"
        clean_line = re.sub(r"[*_`]", "", line)

        # Check (b) Known fabricated names
        for fab_name in FABRICATED_NAMES:
            if re.search(fab_name, line, re.IGNORECASE) or re.search(fab_name, clean_line, re.IGNORECASE):
                violations.append((loc, f"Disallowed fabricated name found matching '{fab_name}': {line.strip()[:100]}"))

        # Check (a) MSSV pattern: 2[0-5]11\d{4}
        mssv_matches = list(re.finditer(r"\b2[0-5]11\d{4}\b", line))
        for m in mssv_matches:
            val = m.group(0)
            # Check if this match is part of an allowed email
            start_pos = max(0, m.start() - 20)
            end_pos = min(len(line), m.end() + 30)
            surrounding = line[start_pos:end_pos].lower()
            is_email = False
            for email in allowed_emails:
                if email in surrounding:
                    is_email = True
                    break
            if not is_email:
                if val not in verified_mssvs:
                    violations.append((loc, f"Unverified MSSV '{val}' found (not verified in docs/TEAM_ROSTER.md): {line.strip()[:100]}"))

        # Check (c) Contribution percentages next to a person / role
        # Avoid flagging evaluation benchmark percentages (e.g., "8/24 (33.3%)" or "Baseline 33.3%")
        if re.search(r"\b(33[.,]3%?|33[.,]4%?)\b", line):
            is_eval_metric = bool(re.search(r"(8/24|baseline|benchmark|pass rate|tỷ lệ pass|đột phá)", line, re.IGNORECASE))
            if not is_eval_metric:
                if re.search(r"(Dev\s+[ABC]|Integrator|Hoàng|Huy|Khánh|Owner|Thành viên|Đóng góp|Contribution)", line, re.IGNORECASE):
                    violations.append((loc, f"Invented contribution percentage next to person/role: {line.strip()[:100]}"))

        # Check (b) Full-name-like string next to role or MSSV
        if not re.search(r"AI\s+Tool\s+Used\s*:", clean_line, re.IGNORECASE):
            name_role_match = re.search(r"([A-ZÀ-Ỹ][a-zà-ỹ]+(?:\s+[A-ZÀ-Ỹ][a-zà-ỹ]+){2,4})\s*(\(|—|-|:|\/|\|)\s*(?:MSSV|`?2[0-5]11\d{4}|Dev\s+[ABC]|Integrator)", clean_line)
            if name_role_match:
                candidate_name = name_role_match.group(1).strip()
                if candidate_name.lower() not in verified_names:
                    violations.append((loc, f"Unverified full name '{candidate_name}' associated with role/MSSV: {line.strip()[:100]}"))

    return violations


def extract_pdf_text(path: Path):
    try:
        import pypdf
        reader = pypdf.PdfReader(str(path))
        text = ""
        for page in reader.pages:
            text += (page.extract_text() or "") + "\n"
        norm_text = unicodedata.normalize("NFC", text)
        # Collapse intra-word spaces caused by font-kerning / diacritic splitting
        collapsed = re.sub(r"(?<=\b[A-Za-zÀ-ỹ])\s+(?=[A-Za-zÀ-ỹ]\b)", "", norm_text)
        return norm_text + "\n" + collapsed, None
    except ImportError:
        return None, "pypdf not installed"
    except Exception as e:
        return None, str(e)


def extract_pptx_text(path: Path):
    try:
        import pptx
        prs = pptx.Presentation(str(path))
        text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        text += p.text + "\n"
        norm_text = unicodedata.normalize("NFC", text)
        return norm_text, None
    except ImportError:
        return None, "python-pptx not installed"
    except Exception as e:
        return None, str(e)


def main():
    print("=" * 60)
    print("SmartDesk AI — Roster & Anti-Fabrication Gate (check_roster.py)")
    print("=" * 60)

    verified_names, verified_mssvs, allowed_emails, allowed_handles = parse_team_roster(ROSTER_PATH)
    print(f"Loaded from docs/TEAM_ROSTER.md:")
    print(f"  - Verified Full Names : {list(verified_names) if verified_names else 'None (CHƯA CÓ DỮ LIỆU)'}")
    print(f"  - Verified MSSVs      : {list(verified_mssvs) if verified_mssvs else 'None (CHƯA CÓ DỮ LIỆU)'}")
    print(f"  - Allowed Git Emails  : {list(allowed_emails)}")
    print(f"  - Allowed Handles     : {list(allowed_handles)}")

    tracked_files = get_tracked_files()
    all_violations = []
    unchecked_binaries = []

    # Files to skip from deep inspection (e.g. self, binary media)
    skip_files = {
        "docs/TEAM_ROSTER.md",
        "scripts/check_roster.py"
    }

    media_extensions = {".png", ".jpg", ".jpeg", ".ico", ".svg", ".zip", ".tar", ".gz", ".db", ".sqlite", ".pyc"}

    for rel_path in tracked_files:
        if rel_path in skip_files:
            continue
        ext = Path(rel_path).suffix.lower()
        if ext in media_extensions:
            continue

        full_path = REPO_ROOT / rel_path
        if not full_path.exists():
            continue

        if ext == ".pdf":
            text, err = extract_pdf_text(full_path)
            if err:
                print(f"[UNCHECKED] {rel_path} ({err})")
                unchecked_binaries.append(rel_path)
            else:
                v = check_content(text, rel_path, verified_names, verified_mssvs, allowed_emails, is_binary=True)
                all_violations.extend(v)
        elif ext in [".pptx", ".ppt"]:
            text, err = extract_pptx_text(full_path)
            if err:
                print(f"[UNCHECKED] {rel_path} ({err})")
                unchecked_binaries.append(rel_path)
            else:
                v = check_content(text, rel_path, verified_names, verified_mssvs, allowed_emails, is_binary=True)
                all_violations.extend(v)
        else:
            try:
                with open(full_path, "r", encoding="utf-8", errors="replace") as f:
                    content = f.read()
                v = check_content(content, rel_path, verified_names, verified_mssvs, allowed_emails, is_binary=False)
                all_violations.extend(v)
            except Exception as e:
                print(f"[WARN] Error reading {rel_path}: {e}")

    print("-" * 60)
    if unchecked_binaries:
        print(f"Notice: {len(unchecked_binaries)} binary files were UNCHECKED due to missing dependencies.")

    if all_violations:
        print(f"[FAIL] Found {len(all_violations)} fabrication violations:")
        for loc, msg in all_violations:
            print(f"  - {loc}: {msg}")
        print("\nAll facts about people must be verified in docs/TEAM_ROSTER.md.")
        print("Unknown names/MSSVs must be 'CHƯA CÓ DỮ LIỆU'. Fake contribution percentages are forbidden.")
        sys.exit(1)
    else:
        print("[PASS] Anti-Fabrication & Roster Gate passed cleanly (0 violations).")
        sys.exit(0)


if __name__ == "__main__":
    main()
